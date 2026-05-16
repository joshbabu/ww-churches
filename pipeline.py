#!/usr/bin/env python3
"""
Word & Worship — Song Pipeline
================================
Converts a Word (.docx) or Excel (.xlsx) file into songs.json
and optionally pushes to GitHub.

Usage:
  python3 pipeline.py --input songs.docx --church bethel-assembly
  python3 pipeline.py --input songs.xlsx --church bethel-assembly
  python3 pipeline.py --input songs.docx --church bethel-assembly --push --token ghp_xxx

Output:
  churches/<church-id>/songs.json
"""

import argparse
import json
import os
import re
import sys
import base64
import urllib.request
import urllib.error

# ── Dependencies check ────────────────────────────────────────────────────────
try:
    import docx
except ImportError:
    print("Installing python-docx...")
    os.system("pip install python-docx --break-system-packages -q")
    import docx

try:
    from openpyxl import load_workbook
except ImportError:
    print("Installing openpyxl...")
    os.system("pip install openpyxl --break-system-packages -q")
    from openpyxl import load_workbook


# ── Helpers ───────────────────────────────────────────────────────────────────

def has_te(text):
    return bool(re.search(r'[\u0C00-\u0C7F]', text))

def get_ref_word(lines):
    for line in lines:
        m = re.search(r'\|\|\s*(.+?)\s*\|\|', line)
        if m:
            return m.group(1).strip()
    return None

def remove_ref(line):
    return re.sub(r'\s*\|\|\s*.+?\s*\|\|\s*$', '', line).strip()

def has_ref(line):
    return bool(re.search(r'\|\|', line))

def to_en_title(te):
    """Simple Telugu to English transliteration for title"""
    mapping = [
        ('అ','A'),('ఆ','Aa'),('ఇ','I'),('ఈ','Ee'),('ఉ','U'),('ఊ','Oo'),
        ('ఎ','E'),('ఏ','Ae'),('ఐ','Ai'),('ఒ','O'),('ఓ','Oh'),('ఔ','Au'),
        ('క','K'),('ఖ','Kh'),('గ','G'),('ఘ','Gh'),('చ','Ch'),('జ','J'),
        ('డ','D'),('త','Th'),('థ','Th'),('ద','D'),('ధ','Dh'),('న','N'),
        ('ప','P'),('ఫ','Ph'),('బ','B'),('భ','Bh'),('మ','M'),('య','Y'),
        ('ర','R'),('ల','L'),('వ','V'),('శ','Sh'),('ష','Sh'),('స','S'),
        ('హ','H'),('ళ','L'),('ట','T'),('ణ','N'),('ా','a'),('ి','i'),
        ('ీ','ee'),('ు','u'),('ూ','oo'),('ె','e'),('ే','ae'),('ై','ai'),
        ('ొ','o'),('ో','oh'),('ౌ','au'),('ం','m'),('ః','h'),('్',''),
    ]
    result = te
    for tc, ec in mapping:
        result = result.replace(tc, ec)
    return ''.join(c if ord(c) < 128 else '' for c in result).strip()


# ── Parse Telugu lines ────────────────────────────────────────────────────────

def parse_te_lines(lines):
    ref_word = get_ref_word(lines)
    blocks = []
    current_type = 'chorus'
    current_label = 'పల్లవి'
    current_lines = []

    for line in lines:
        ap_m = re.match(r'^అ\.ప\s*:?\s*(.*)', line)
        if ap_m:
            content = ap_m.group(1).strip()
            if content:
                current_lines.append(content)
            continue

        vm = re.match(r'^(\d+)\)(.*)', line)
        if vm:
            if current_lines:
                text = '\n'.join(l for l in current_lines if l.strip())
                if text:
                    blocks.append({'type': current_type, 'label': current_label, 'text': text})
            current_type = 'verse'
            current_label = f'{vm.group(1)})'
            rest = vm.group(2).strip()
            if has_ref(rest):
                clean = remove_ref(rest)
                blocks.append({'type': 'verse', 'label': current_label, 'text': clean})
                if ref_word:
                    blocks.append({'type': 'chorusRef', 'text': f'|| {ref_word} ||'})
                current_lines = []
            else:
                current_lines = [rest] if rest else []
            continue

        if has_ref(line):
            clean = remove_ref(line)
            if clean:
                current_lines.append(clean)
            if current_lines:
                text = '\n'.join(l for l in current_lines if l.strip())
                if text:
                    blocks.append({'type': current_type, 'label': current_label, 'text': text})
            if ref_word:
                blocks.append({'type': 'chorusRef', 'text': f'|| {ref_word} ||'})
            current_lines = []
            continue

        current_lines.append(line)

    if current_lines:
        text = '\n'.join(l for l in current_lines if l.strip())
        if text:
            blocks.append({'type': current_type, 'label': current_label, 'text': text})

    return blocks


# ── Parse English lines ───────────────────────────────────────────────────────

def parse_en_lines(en_lines, te_blocks):
    ref_word = get_ref_word(en_lines)
    te_chorus = [b for b in te_blocks if b['type'] == 'chorus']
    te_verses = [b for b in te_blocks if b['type'] == 'verse']
    n_te_chorus_lines = len(te_chorus[0]['text'].split('\n')) if te_chorus else 1
    n_te_verses = len(te_verses)

    has_ap = any(re.match(r'^A\.Pa\s*:?', l) for l in en_lines)
    n_refs = sum(1 for l in en_lines if has_ref(l))
    blocks = []

    if has_ap:
        try:
            ap_idx = next(i for i, l in enumerate(en_lines) if re.match(r'^A\.Pa\s*:?', l))
        except StopIteration:
            ap_idx = 0
        title_lines = [l for l in en_lines[:ap_idx] if l.strip()]
        ap_content = re.sub(r'^A\.Pa\s*:?\s*', '', en_lines[ap_idx]).strip()
        rest_lines = en_lines[ap_idx + 1:]
        chorus_lines = title_lines + ([ap_content] if ap_content else [])

        sections = []
        current = []
        for line in rest_lines:
            if has_ref(line):
                clean = remove_ref(line)
                if clean:
                    current.append(clean)
                sections.append(list(current))
                current = []
            else:
                current.append(line)
        if current:
            sections.append(list(current))

        if sections:
            if len(chorus_lines) >= n_te_chorus_lines:
                blocks.append({'type': 'chorus', 'label': 'Chorus', 'text': '\n'.join(chorus_lines)})
                if sections[0]:
                    blocks.append({'type': 'verse', 'label': '1)', 'text': '\n'.join(sections[0])})
                    if ref_word:
                        blocks.append({'type': 'chorusRef', 'text': f'|| {ref_word} ||'})
                for vi, sec in enumerate(sections[1:], 2):
                    if sec:
                        blocks.append({'type': 'verse', 'label': f'{vi})', 'text': '\n'.join(sec)})
                        if ref_word:
                            blocks.append({'type': 'chorusRef', 'text': f'|| {ref_word} ||'})
            else:
                combined = chorus_lines + sections[0]
                blocks.append({'type': 'chorus', 'label': 'Chorus', 'text': '\n'.join(combined)})
                if ref_word:
                    blocks.append({'type': 'chorusRef', 'text': f'|| {ref_word} ||'})
                for vi, sec in enumerate(sections[1:], 2):
                    if sec:
                        blocks.append({'type': 'verse', 'label': f'{vi})', 'text': '\n'.join(sec)})
                        if ref_word:
                            blocks.append({'type': 'chorusRef', 'text': f'|| {ref_word} ||'})
        else:
            blocks.append({'type': 'chorus', 'label': 'Chorus', 'text': '\n'.join(chorus_lines)})
        return blocks

    if n_refs == 0:
        non_empty = [l for l in en_lines if l.strip()]
        idx = min(n_te_chorus_lines, len(non_empty))
        blocks.append({'type': 'chorus', 'label': 'Chorus', 'text': '\n'.join(non_empty[:idx])})
        for vi, tv in enumerate(te_verses, 1):
            count = len(tv['text'].split('\n'))
            chunk = non_empty[idx:idx + count]
            if chunk:
                blocks.append({'type': 'verse', 'label': f'{vi})', 'text': '\n'.join(chunk)})
            idx += count
        return blocks

    # Has refs
    sections = []
    current = []
    for line in en_lines:
        if has_ref(line):
            clean = remove_ref(line)
            if clean:
                current.append(clean)
            sections.append(list(current))
            current = []
        else:
            current.append(line)
    if current:
        sections.append(list(current))

    if n_refs == n_te_verses + 1:
        if sections[0]:
            blocks.append({'type': 'chorus', 'label': 'Chorus', 'text': '\n'.join(s for s in sections[0] if s)})
        if ref_word:
            blocks.append({'type': 'chorusRef', 'text': f'|| {ref_word} ||'})
        for vi, sec in enumerate(sections[1:n_te_verses + 1], 1):
            text = '\n'.join(s for s in sec if s)
            if text:
                blocks.append({'type': 'verse', 'label': f'{vi})', 'text': text})
                if ref_word:
                    blocks.append({'type': 'chorusRef', 'text': f'|| {ref_word} ||'})
    else:
        sec0 = [l for l in sections[0] if l.strip()]
        chorus_part = sec0[:n_te_chorus_lines]
        verse1_part = sec0[n_te_chorus_lines:]
        if chorus_part:
            blocks.append({'type': 'chorus', 'label': 'Chorus', 'text': '\n'.join(chorus_part)})
        if verse1_part:
            blocks.append({'type': 'verse', 'label': '1)', 'text': '\n'.join(verse1_part)})
            if ref_word:
                blocks.append({'type': 'chorusRef', 'text': f'|| {ref_word} ||'})
        for vi, sec in enumerate(sections[1:n_te_verses], 2):
            text = '\n'.join(s for s in sec if s)
            if text:
                blocks.append({'type': 'verse', 'label': f'{vi})', 'text': text})
                if ref_word:
                    blocks.append({'type': 'chorusRef', 'text': f'|| {ref_word} ||'})
        if len(sections) > n_te_verses:
            last = '\n'.join(s for s in sections[n_te_verses] if s)
            if last:
                blocks.append({'type': 'verse', 'label': f'{n_te_verses})', 'text': last})

    return blocks if blocks else [{'type': 'chorus', 'label': 'Chorus', 'text': '\n'.join(en_lines)}]


# ── Extract lines from docx ───────────────────────────────────────────────────

def extract_docx_lines(filepath):
    """Extract all paragraph lines from a Word file (handles nested tables)"""
    doc = docx.Document(filepath)

    def get_paras(element):
        paras = []
        for child in element:
            tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
            if tag == 'p':
                paras.append(child)
            elif tag in ('tbl', 'tr', 'tc', 'body'):
                paras.extend(get_paras(child))
        return paras

    def para_to_text(para_elem):
        lines = []
        current = []
        for elem in para_elem.iter():
            tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag
            if tag == 't' and elem.text:
                current.append(elem.text)
            elif tag == 'br':
                lines.append(''.join(current))
                current = []
        if current:
            lines.append(''.join(current))
        return [l.strip() for l in lines if l.strip()]

    all_paras = get_paras(doc.element.body)
    output = []
    for para in all_paras:
        for line in para_to_text(para):
            output.append(line)
    return output


# ── Parse docx into songs ─────────────────────────────────────────────────────

def parse_docx(filepath):
    lines = extract_docx_lines(filepath)
    songs_raw = {}
    current_num = None
    current_lines = []
    is_te = True

    for line in lines:
        m = re.match(r'^Song (\d{3,4})$', line.strip())
        if m:
            num = m.group(1)
            if current_num:
                if current_num not in songs_raw:
                    songs_raw[current_num] = {'te': [], 'en': []}
                if is_te:
                    songs_raw[current_num]['te'] = current_lines
                else:
                    songs_raw[current_num]['en'] = current_lines
            if num == current_num:
                is_te = False
            else:
                current_num = num
                is_te = True
            current_lines = []
        else:
            current_lines.append(line)

    if current_num:
        if current_num not in songs_raw:
            songs_raw[current_num] = {'te': [], 'en': []}
        if is_te:
            songs_raw[current_num]['te'] = current_lines
        else:
            songs_raw[current_num]['en'] = current_lines

    songs = []
    for num in sorted(songs_raw.keys(), key=int):
        te_lines = songs_raw[num].get('te', [])
        en_lines = songs_raw[num].get('en', [])
        if te_lines and not has_te(' '.join(te_lines[:3])):
            te_lines, en_lines = en_lines, te_lines

        te_blocks = parse_te_lines(te_lines)
        en_blocks = parse_en_lines(en_lines, te_blocks)

        te_title = te_lines[0] if te_lines else ''
        en_title = en_lines[0] if en_lines else to_en_title(te_title)

        songs.append({
            'no': int(num),
            'oldNo': int(num),
            'catId': 11,
            'te': te_title,
            'en': en_title,
            'author': '',
            'authorEn': '',
            'lyricsTE': te_blocks,
            'lyricsEN': en_blocks,
        })

    print(f"  Parsed {len(songs)} songs from Word file")
    return songs


# ── Parse Excel into songs ────────────────────────────────────────────────────

def parse_excel(filepath):
    """
    Expected Excel columns:
    no | te (title) | en (title) | author | authorEn | catId | chorus_te | chorus_en | verse1_te | verse1_en | ...
    
    OR simpler format:
    no | te | en | author | chorus | verse1 | verse2 | verse3
    """
    wb = load_workbook(filepath, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).lower().strip() if h else '' for h in rows[0]]
    print(f"  Excel headers: {headers}")

    songs = []
    for row in rows[1:]:
        if not row[0]:
            continue
        r = dict(zip(headers, row))

        no = int(float(str(r.get('no', 0)))) if r.get('no') else 0
        te = str(r.get('te', '') or r.get('title_te', '') or '').strip()
        en = str(r.get('en', '') or r.get('title_en', '') or to_en_title(te)).strip()
        author = str(r.get('author', '') or '').strip()
        author_en = str(r.get('authoren', '') or r.get('author_en', '') or '').strip()
        cat_id = int(float(str(r.get('catid', 11) or 11)))

        # Build lyrics blocks from columns
        lyrics_te = []
        lyrics_en = []

        chorus_te = str(r.get('chorus_te', '') or r.get('chorus', '') or '').strip()
        chorus_en = str(r.get('chorus_en', '') or '').strip()
        if chorus_te:
            lyrics_te.append({'type': 'chorus', 'label': 'పల్లవి', 'text': chorus_te})
        if chorus_en:
            lyrics_en.append({'type': 'chorus', 'label': 'Chorus', 'text': chorus_en})

        vi = 1
        while True:
            vte = str(r.get(f'verse{vi}_te', '') or r.get(f'verse{vi}', '') or '').strip()
            ven = str(r.get(f'verse{vi}_en', '') or '').strip()
            if not vte and not ven:
                break
            if vte:
                lyrics_te.append({'type': 'verse', 'label': f'{vi})', 'text': vte})
            if ven:
                lyrics_en.append({'type': 'verse', 'label': f'{vi})', 'text': ven})
            vi += 1

        if not lyrics_te:
            lyrics_te = [{'type': 'chorus', 'label': 'పల్లవి', 'text': 'Lyrics coming soon...'}]
        if not lyrics_en:
            lyrics_en = [{'type': 'chorus', 'label': 'Chorus', 'text': 'Lyrics coming soon...'}]

        songs.append({
            'no': no,
            'oldNo': no,
            'catId': cat_id,
            'te': te,
            'en': en,
            'author': author,
            'authorEn': author_en,
            'lyricsTE': lyrics_te,
            'lyricsEN': lyrics_en,
        })

    print(f"  Parsed {len(songs)} songs from Excel file")
    return songs


# ── Push to GitHub ────────────────────────────────────────────────────────────

def push_to_github(songs, church_id, token, repo='joshbabu/ww-churches'):
    path = f'churches/{church_id}/songs.json'
    owner, repo_name = repo.split('/')
    content = json.dumps(songs, ensure_ascii=False, indent=2)
    encoded = base64.b64encode(content.encode('utf-8')).decode()

    # Get current SHA if file exists
    sha = None
    try:
        req = urllib.request.Request(
            f'https://api.github.com/repos/{owner}/{repo_name}/contents/{path}',
            headers={'Authorization': f'token {token}', 'Accept': 'application/vnd.github+json'}
        )
        with urllib.request.urlopen(req) as r:
            sha = json.loads(r.read())['sha']
    except urllib.error.HTTPError:
        pass

    body = json.dumps({
        'message': f'Update songs for {church_id} via pipeline',
        'content': encoded,
        **(({'sha': sha}) if sha else {})
    }).encode()

    req = urllib.request.Request(
        f'https://api.github.com/repos/{owner}/{repo_name}/contents/{path}',
        data=body,
        method='PUT',
        headers={
            'Authorization': f'token {token}',
            'Content-Type': 'application/json',
            'Accept': 'application/vnd.github+json'
        }
    )
    with urllib.request.urlopen(req) as r:
        result = json.loads(r.read())
        print(f"  ✅ Pushed to GitHub: {result['content']['html_url']}")


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description='Word & Worship Song Pipeline')
    parser.add_argument('--input', required=True, help='Input file (.docx or .xlsx)')
    parser.add_argument('--church', required=True, help='Church ID (e.g. bethel-assembly)')
    parser.add_argument('--push', action='store_true', help='Push to GitHub')
    parser.add_argument('--token', help='GitHub personal access token')
    parser.add_argument('--repo', default='joshbabu/ww-churches', help='GitHub repo')
    parser.add_argument('--output', help='Output path (default: churches/<id>/songs.json)')
    args = parser.parse_args()

    print(f"\n🎵 Word & Worship Song Pipeline")
    print(f"   Input:  {args.input}")
    print(f"   Church: {args.church}")
    print()

    # Parse input
    ext = os.path.splitext(args.input)[1].lower()
    if ext == '.docx':
        print("📖 Parsing Word file...")
        songs = parse_docx(args.input)
    elif ext in ('.xlsx', '.xls'):
        print("📊 Parsing Excel file...")
        songs = parse_excel(args.input)
    else:
        print(f"❌ Unsupported file type: {ext}")
        sys.exit(1)

    # Save locally
    output_path = args.output or f'churches/{args.church}/songs.json'
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(songs, f, ensure_ascii=False, indent=2)
    print(f"\n✅ Saved {len(songs)} songs to {output_path}")

    # Push to GitHub
    if args.push:
        if not args.token:
            print("❌ --token required when using --push")
            sys.exit(1)
        print(f"\n⬆️  Pushing to GitHub ({args.repo})...")
        push_to_github(songs, args.church, args.token, args.repo)

    print(f"\n🎉 Done! {len(songs)} songs ready.")
    print(f"   Next: update config.json and deploy the app.")


if __name__ == '__main__':
    main()
